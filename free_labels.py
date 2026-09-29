"""Excel import, A4 label geometry and shared preview/print rendering."""

from dataclasses import dataclass
from datetime import date, datetime, time
from functools import lru_cache
import math
import os
from pathlib import Path
import re

from PIL import Image, ImageDraw, ImageFont

PAGE_W, PAGE_H = 210.0, 297.0
MAX_ROWS = 50000
MAX_TEXT = 4000
RENDER_DPI = 300


@dataclass(frozen=True)
class LabelFormat:
    code: str
    width: float
    height: float
    columns: int
    rows: int
    left: float
    top: float
    gap_x: float = 0.0
    gap_y: float = 0.0

    @property
    def per_sheet(self):
        return self.columns * self.rows

    @property
    def description(self):
        return (f"{self.code} · {self.width:g} × {self.height:g} mm · "
                f"{self.per_sheet}/Bogen").replace(".", ",")

    def validate(self):
        values = (self.width, self.height, self.left, self.top, self.gap_x, self.gap_y)
        if not all(math.isfinite(v) for v in values):
            raise ValueError("Bitte gültige Maße in mm eingeben.")
        if self.width < 5 or self.height < 5:
            raise ValueError("Etiketten müssen mindestens 5 × 5 mm groß sein.")
        if any(v < 0 for v in (self.left, self.top, self.gap_x, self.gap_y)):
            raise ValueError("Ränder und Abstände dürfen nicht negativ sein.")
        if not (1 <= self.columns <= 40 and 1 <= self.rows <= 59):
            raise ValueError("Bitte 1–40 Spalten und 1–59 Reihen wählen.")
        right = self.left + self.columns * self.width + (self.columns - 1) * self.gap_x
        bottom = self.top + self.rows * self.height + (self.rows - 1) * self.gap_y
        if right > PAGE_W + 0.01 or bottom > PAGE_H + 0.01:
            raise ValueError("Das eingestellte Raster passt nicht auf einen A4-Bogen.")
        return self

    def rect(self, position):
        if not 0 <= position < self.per_sheet:
            raise ValueError("Etikettenposition außerhalb des Bogens.")
        row, col = divmod(position, self.columns)
        return (self.left + col * (self.width + self.gap_x),
                self.top + row * (self.height + self.gap_y), self.width, self.height)


# Nominal dimensions: Avery Zweckform's product/template pages.
# Grid pitch and margins cross-checked against LibreOffice's label catalogue.
# See FORMAT_SOURCES.md for sources and the small nominal rounding differences.
FORMATS = [
    LabelFormat("L4761", 192, 61, 1, 4, 9, 26.5),
    LabelFormat("L7159", 63.5, 33.9, 3, 8, 7.21, 12.9, 2.54),
    LabelFormat("L7160", 63.5, 38.1, 3, 7, 7.21, 15.15, 2.54),
    LabelFormat("L7161", 63.5, 46.6, 3, 6, 7.21, 8.7, 2.54),
    LabelFormat("L7162", 99.1, 33.9, 2, 8, 4.63, 12.9, 2.54),
    LabelFormat("L7163", 99.1, 38.1, 2, 7, 4.63, 15.15, 2.54),
    LabelFormat("L7164", 63.5, 72, 3, 4, 7.21, 4.5, 2.54),
    LabelFormat("L7165", 99.1, 67.7, 2, 4, 4.63, 13.1, 2.54),
    LabelFormat("L7166", 99.1, 93.1, 2, 3, 4.63, 8.85, 2.54),
    LabelFormat("L7167", 199.6, 289.1, 1, 1, 5.2, 3.95),
    LabelFormat("L7168", 199.6, 143.5, 1, 2, 5.2, 5),
    LabelFormat("L7169", 99.1, 139, 2, 2, 4.63, 9.5, 2.54),
    LabelFormat("3666", 38, 21.2, 5, 13, 10, 10.7),
]
FORMATS_BY_CODE = {f.code: f for f in FORMATS}


@dataclass(frozen=True)
class ExcelRow:
    number: int
    text: str
    error: str = ""


def cell_text(value, number_format="General"):
    """Preserve literal text, line breaks and simple numeric zero masks."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "WAHR" if value else "FALSCH"
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y %H:%M:%S" if value.time() != time() else "%d.%m.%Y")
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    if isinstance(value, (int, float)):
        if not math.isfinite(value):
            raise ValueError("Ungültiger Zahlenwert.")
        mask = number_format.split(";")[0]
        if value == int(value):
            if re.fullmatch(r"0+", mask):
                sign = "-" if value < 0 else ""
                return sign + str(abs(int(value))).zfill(len(mask))
            return str(int(value))
        return str(value).replace(".", ",")
    return str(value).replace("\r\n", "\n").replace("\r", "\n").replace("\t", "    ").strip()


def _row(number, value, number_format="General", error=""):
    text = cell_text(value, number_format)
    if len(text) > MAX_TEXT:
        error = f"Zelltext ist zu lang (maximal {MAX_TEXT} Zeichen)."
    return ExcelRow(number, text, error)


def read_excel(path):
    """Read column A of all sheets; never execute macros or calculate formulas.

    Return rows with source row numbers, including blanks and per-cell errors.
    Keeping the raw rows makes header/blank options reversible without rereading.
    """
    suffix = Path(path).suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        import openpyxl
        result = {}
        # The formula view detects missing cached results instead of dropping them.
        with open(path, "rb") as values_file, open(path, "rb") as formulas_file:
            values = openpyxl.load_workbook(values_file, read_only=True, data_only=True)
            try:
                formulas = openpyxl.load_workbook(formulas_file, read_only=True, data_only=False)
                try:
                    count = 0
                    for sheet in values:
                        if sheet.max_row and sheet.max_row > MAX_ROWS:
                            raise ValueError(f"„{sheet.title}“ hat mehr als {MAX_ROWS} Zeilen. "
                                             "Bitte den benötigten Bereich in eine neue Datei kopieren.")
                        rows = []
                        raw_sheet = formulas[sheet.title]
                        for number, (cells, raw_cells) in enumerate(zip(
                                sheet.iter_rows(max_col=1), raw_sheet.iter_rows(max_col=1)), 1):
                            count += 1
                            if count > MAX_ROWS:
                                raise ValueError(f"Die Datei darf insgesamt höchstens {MAX_ROWS} Zeilen enthalten.")
                            cell, raw = cells[0], raw_cells[0]
                            error = ""
                            if raw.data_type == "f" and cell.value is None and cell.data_type != "str":
                                error = "Formel ohne gespeichertes Ergebnis: in Excel neu berechnen und speichern oder Werte einfügen."
                            elif cell.data_type == "e":
                                error = f"Excel-Fehler: {cell.value}"
                            rows.append(_row(number, cell.value, cell.number_format, error))
                        result[sheet.title] = rows
                finally:
                    formulas.close()
            finally:
                values.close()
        return result
    if suffix == ".xls":
        import xlrd
        book = xlrd.open_workbook(path, on_demand=True, formatting_info=True)
        try:
            result, count = {}, 0
            for sheet in book.sheets():
                count += sheet.nrows
                if count > MAX_ROWS:
                    raise ValueError(f"Die Datei darf insgesamt höchstens {MAX_ROWS} Zeilen enthalten.")
                rows = []
                for i in range(sheet.nrows):
                    if not sheet.ncols:
                        break
                    cell = sheet.cell(i, 0)
                    value, error = cell.value, ""
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate_as_datetime(value, book.datemode)
                        if cell.value < 1:
                            value = value.time()
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        value = bool(value)
                    elif cell.ctype == xlrd.XL_CELL_ERROR:
                        error = f"Excel-Fehler: {xlrd.error_text_from_code.get(value, value)}"
                    fmt = book.format_map[book.xf_list[cell.xf_index].format_key].format_str
                    rows.append(_row(i + 1, value, fmt, error))
                result[sheet.name] = rows
            return result
        finally:
            book.release_resources()
    raise ValueError("Bitte eine Excel-Datei (.xlsx, .xlsm oder .xls) auswählen.")


def excel_items(rows, skip_header=False, keep_blanks=False):
    rows = [row for row in rows if not (skip_header and row.number == 1)]
    # Ignore unused trailing cells (often only formatted in Excel).
    while rows and not rows[-1].text and not rows[-1].error:
        rows.pop()
    errors = [f"A{r.number}: {r.error}" for r in rows if r.error]
    if errors:
        raise ValueError("\n".join(errors[:5]) + ("\n…" if len(errors) > 5 else ""))
    return [r.text or None for r in rows if r.text or keep_blanks]


def paginate(items, start, per_sheet):
    if not 1 <= start <= per_sheet:
        raise ValueError(f"Die Startposition muss zwischen 1 und {per_sheet} liegen.")
    slots = [None] * (start - 1) + list(items) if items else []
    return [slots[i:i + per_sheet] + [None] * (per_sheet - len(slots[i:i + per_sheet]))
            for i in range(0, max(1, len(slots)), per_sheet)]


@lru_cache(maxsize=128)
def label_font(size):
    candidates = [str(Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts" / "arialbd.ttf"),
                  "DejaVuSans-Bold.ttf"]
    for candidate in candidates:
        try:
            return ImageFont.truetype(candidate, size)
        except OSError:
            pass
    return ImageFont.load_default(size=size)


def wrap_text(text, font, width):
    """Wrap at spaces; split long identifiers only when they cannot fit."""
    lines = []
    for paragraph in text.split("\n"):
        line = ""
        for word in paragraph.split():
            candidate = f"{line} {word}" if line else word
            if font.getlength(candidate) <= width:
                line = candidate
                continue
            if line:
                lines.append(line)
                line = ""
            while font.getlength(word) > width and len(word) > 1:
                lo, hi = 1, len(word)
                while lo < hi:
                    middle = (lo + hi + 1) // 2
                    if font.getlength(word[:middle]) <= width:
                        lo = middle
                    else:
                        hi = middle - 1
                lines.append(word[:lo])
                word = word[lo:]
            line = word
        lines.append(line)
    return lines


def fitted_text(text, width, height, max_size):
    draw = ImageDraw.Draw(Image.new("RGB", (1, 1)))
    lo, hi, best = 1, max(1, max_size), None
    while lo <= hi:
        size = (lo + hi) // 2
        font = label_font(size)
        wrapped = "\n".join(wrap_text(text, font, width))
        spacing = max(1, size // 5)
        box = draw.multiline_textbbox((0, 0), wrapped, font=font, spacing=spacing, align="center")
        if box[2] - box[0] <= width and box[3] - box[1] <= height:
            best = font, wrapped, spacing, box
            lo = size + 1
        else:
            hi = size - 1
    if best is None:
        raise ValueError("Der Text passt nicht in das Etikett. Bitte Text kürzen oder größeres Format wählen.")
    return best


def render_sheet(sheet, layout, font_mm=7, off_x=0, off_y=0, border=False, frames=False):
    """Single 300-dpi renderer used for both preview and Windows printing."""
    layout.validate()
    scale = RENDER_DPI / 25.4
    image = Image.new("RGB", (round(PAGE_W * scale), round(PAGE_H * scale)), "white")
    draw = ImageDraw.Draw(image)
    for position in range(layout.per_sheet):
        left, top, width, height = layout.rect(position)
        x, y = round((left + off_x) * scale), round((top + off_y) * scale)
        w, h = round(width * scale), round(height * scale)
        if frames:
            draw.rectangle((x, y, x + w - 1, y + h - 1), outline="#777777", width=2)
        item = sheet[position] if position < len(sheet) else None
        if item is None:
            continue
        # Render inside a separate label image so text never spills into a neighbour.
        label = Image.new("RGB", (w, h), "white")
        pen = ImageDraw.Draw(label)
        inset = max(2, round(min(2, width / 8, height / 8) * scale))
        if border:
            b = max(1, round(1 * scale))
            pen.rectangle((b, b, w - b - 1, h - b - 1), outline="black", width=max(1, round(.25 * scale)))
        font, text, spacing, bbox = fitted_text(item, w - 2 * inset, h - 2 * inset,
                                                round(font_mm * scale))
        pen.multiline_text(((w - (bbox[2] - bbox[0])) / 2 - bbox[0],
                            (h - (bbox[3] - bbox[1])) / 2 - bbox[1]),
                           text, font=font, fill="black", spacing=spacing, align="center")
        image.paste(label, (x, y))
        if frames:
            draw.rectangle((x, y, x + w - 1, y + h - 1), outline="#777777", width=2)
    return image


def print_excel_sheets(printer, sheets, layout, options, devmode=None):
    import win32con
    import win32gui
    import win32print
    import win32ui
    from PIL import ImageWin

    handle = win32print.OpenPrinter(printer)
    try:
        dm = devmode or win32print.GetPrinter(handle, 2).get("pDevMode")
        if dm is None:
            raise RuntimeError("Der Drucker liefert keine Papier-Einstellungen.")
        dm.PaperSize = win32con.DMPAPER_A4
        dm.Orientation = win32con.DMORIENT_PORTRAIT
        dm.Scale = 100
        dm.Fields &= ~(win32con.DM_PAPERLENGTH | win32con.DM_PAPERWIDTH)
        dm.Fields |= win32con.DM_PAPERSIZE | win32con.DM_ORIENTATION | win32con.DM_SCALE
        win32print.DocumentProperties(0, handle, printer, dm, dm,
                                     win32con.DM_IN_BUFFER | win32con.DM_OUT_BUFFER)
        dc = win32ui.CreateDCFromHandle(win32gui.CreateDC("WINSPOOL", printer, dm))
    finally:
        win32print.ClosePrinter(handle)
    started = False
    try:
        dx, dy = dc.GetDeviceCaps(win32con.LOGPIXELSX), dc.GetDeviceCaps(win32con.LOGPIXELSY)
        ox, oy = dc.GetDeviceCaps(win32con.PHYSICALOFFSETX), dc.GetDeviceCaps(win32con.PHYSICALOFFSETY)
        paper_w = dc.GetDeviceCaps(win32con.PHYSICALWIDTH) / dx * 25.4
        paper_h = dc.GetDeviceCaps(win32con.PHYSICALHEIGHT) / dy * 25.4
        if abs(paper_w - PAGE_W) > 2 or abs(paper_h - PAGE_H) > 2:
            raise RuntimeError("Bitte im Druckertreiber A4 im Hochformat einstellen.")
        dc.StartDoc("Freie Excel-Etiketten")
        started = True
        for sheet in sheets:
            image = render_sheet(sheet, layout, **options)
            dc.StartPage()
            ImageWin.Dib(image).draw(dc.GetHandleOutput(),
                (-ox, -oy, round(PAGE_W / 25.4 * dx) - ox, round(PAGE_H / 25.4 * dy) - oy))
            dc.EndPage()
        dc.EndDoc()
        started = False
    finally:
        if started:
            dc.AbortDoc()
        dc.DeleteDC()
