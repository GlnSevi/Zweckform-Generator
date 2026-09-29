import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from types import SimpleNamespace
import zipfile

from openpyxl import Workbook
from PIL import ImageChops

from free_labels import (ExcelRow, FORMATS, FORMATS_BY_CODE, LabelFormat, MAX_ROWS,
                         cell_text, excel_items, fitted_text, paginate, read_excel,
                         render_sheet, RENDER_DPI, CURATED, search_formats)


class ExcelTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.path = Path(self.tmp.name) / "labels.xlsx"

    def save(self, values):
        book = Workbook()
        sheet = book.active
        sheet.title = "Etiketten"
        for i, value in enumerate(values, 1):
            sheet.cell(i, 1, value)
        sheet["B1"] = "Diese Spalte darf nicht importiert werden"
        book.save(self.path)
        book.close()

    def test_column_a_order_header_blanks_and_literal_pipe(self):
        self.save(["Überschrift", "00123", None, "Text | 3", "Mehrzeilig\nÖl für Geräte", 0, False])
        rows = read_excel(self.path)["Etiketten"]
        self.assertEqual(excel_items(rows, True, True),
            ["00123", None, "Text | 3", "Mehrzeilig\nÖl für Geräte", "0", "FALSCH"])
        self.assertEqual(excel_items(rows, True),
            ["00123", "Text | 3", "Mehrzeilig\nÖl für Geräte", "0", "FALSCH"])
        self.assertEqual(excel_items(rows)[0], "Überschrift")

    def test_zero_mask_multiple_sheets_and_trailing_formatting(self):
        book = Workbook()
        sheet = book.active
        sheet.title = "Erstes Blatt"
        sheet["A1"] = 42
        sheet["A1"].number_format = "000000"
        sheet["A20"].number_format = "0.00"
        book.create_sheet("Zweites Blatt")["A1"] = "Andere Etiketten"
        book.save(self.path)
        book.close()
        imported = read_excel(self.path)
        self.assertEqual(list(imported), ["Erstes Blatt", "Zweites Blatt"])
        self.assertEqual(excel_items(imported["Erstes Blatt"], keep_blanks=True), ["000042"])
        self.assertEqual(excel_items(imported["Zweites Blatt"]), ["Andere Etiketten"])
        # No dangling read-only workbook handles.
        self.path.unlink()

    def test_uncached_formula_is_reported_at_source_row(self):
        self.save(["Text", "=1+2"])
        with self.assertRaisesRegex(ValueError, "A2: Formel"):
            excel_items(read_excel(self.path)["Etiketten"])

    def test_cached_formula_uses_result_not_formula(self):
        self.save(["=1+2"])
        with zipfile.ZipFile(self.path) as source:
            contents = {name: source.read(name) for name in source.namelist()}
        contents["xl/worksheets/sheet1.xml"] = contents["xl/worksheets/sheet1.xml"].replace(b"<v></v>", b"<v>3</v>")
        with zipfile.ZipFile(self.path, "w") as target:
            for name, data in contents.items():
                target.writestr(name, data)
        self.assertEqual(excel_items(read_excel(self.path)["Etiketten"]), ["3"])

    def test_excel_errors_empty_and_bad_files(self):
        self.save(["#DIV/0!"])
        with self.assertRaisesRegex(ValueError, "A1: Excel-Fehler"):
            excel_items(read_excel(self.path)["Etiketten"])
        self.save([])
        self.assertEqual(excel_items(read_excel(self.path)["Etiketten"]), [])
        self.path.write_bytes(b"not an Excel file")
        with self.assertRaises(zipfile.BadZipFile):
            read_excel(self.path)
        with self.assertRaisesRegex(ValueError, "Excel-Datei"):
            read_excel(self.path.with_suffix(".csv"))

    def test_cached_empty_formula_keeps_blank_position(self):
        self.save(['=""', "Next label"])
        with zipfile.ZipFile(self.path) as source:
            contents = {name: source.read(name) for name in source.namelist()}
        contents["xl/worksheets/sheet1.xml"] = contents["xl/worksheets/sheet1.xml"].replace(
            b'<c r="A1">', b'<c r="A1" t="str">')
        with zipfile.ZipFile(self.path, "w") as target:
            for name, data in contents.items():
                target.writestr(name, data)
        self.assertEqual(excel_items(read_excel(self.path)["Etiketten"], keep_blanks=True),
                         [None, "Next label"])

    def test_oversized_sheet_and_text_are_rejected(self):
        book = Workbook()
        book.active.cell(MAX_ROWS + 1, 1, "last")
        book.save(self.path)
        book.close()
        with self.assertRaisesRegex(ValueError, "Zeilen"):
            read_excel(self.path)
        self.save(["x" * 4001])
        with self.assertRaisesRegex(ValueError, "A1: Zelltext"):
            excel_items(read_excel(self.path)["Etiketten"])

    def test_xls_import(self):
        try:
            import xlwt
        except ImportError:
            self.skipTest("Install requirements-dev.txt to create a legacy XLS fixture")
        path = self.path.with_suffix(".xls")
        book = xlwt.Workbook()
        sheet = book.add_sheet("Altformat")
        sheet.write(0, 0, "Überschrift")
        sheet.write(1, 0, 7, xlwt.easyxf(num_format_str="0000"))
        sheet.write(2, 0, "Öl\nWerkstatt")
        sheet.write(2, 1, "Ignored")
        book.save(str(path))
        self.assertEqual(excel_items(read_excel(path)["Altformat"], True), ["0007", "Öl\nWerkstatt"])


class LayoutTests(unittest.TestCase):
    def test_all_formats_fit_a4_and_use_row_major_order(self):
        for layout in FORMATS:
            with self.subTest(layout.code):
                layout.validate()
                self.assertEqual(layout.rect(0)[:2], (layout.left, layout.top))
                if layout.columns > 1:
                    self.assertAlmostEqual(layout.rect(1)[0], layout.left + layout.width + layout.gap_x)
                    self.assertEqual(layout.rect(1)[1], layout.top)
                if layout.rows > 1:
                    self.assertEqual(layout.rect(layout.columns)[0], layout.left)
                    self.assertAlmostEqual(layout.rect(layout.columns)[1], layout.top + layout.height + layout.gap_y)
                values = [f"Zeile {i}" for i in range(layout.per_sheet * 2 + 1)]
                sheets = paginate(values, layout.per_sheet, layout.per_sheet)
                self.assertEqual(sheets[0][:-1], [None] * (layout.per_sheet - 1))
                self.assertEqual([x for page in sheets for x in page if x is not None], values)
                self.assertEqual(len(sheets), 3)

    def test_blank_and_partial_sheets(self):
        self.assertEqual(paginate(["A", None, "B"], 4, 4),
                         [[None, None, None, "A"], [None, "B", None, None]])
        self.assertEqual(paginate([], 1, 4), [[None] * 4])
        for start in (0, 5):
            with self.assertRaises(ValueError):
                paginate(["A"], start, 4)

    def test_invalid_custom_geometry(self):
        for kwargs in ({"width": 300}, {"left": -1}, {"gap_x": -1},
                       {"height": float("nan")}, {"rows": 0}, {"width": 4}):
            values = dict(code="Custom", width=63.5, height=38.1, columns=3, rows=7,
                          left=7.21, top=15.15, gap_x=2.54)
            values.update(kwargs)
            with self.assertRaises(ValueError):
                LabelFormat(**values).validate()

    def test_text_fits_with_newlines_and_long_identifiers(self):
        for text in ["Mehrzeiliger\nText mit Umlauten ÄÖÜ", "LANGEARTIKELNUMMER" * 30, "A " * 300]:
            font, wrapped, spacing, bbox = fitted_text(text, 400, 150, 100)
            self.assertLessEqual(bbox[2] - bbox[0], 400)
            self.assertLessEqual(bbox[3] - bbox[1], 150)
            self.assertEqual("".join(wrapped.split()), "".join(text.split()))

    def test_render_leaves_blank_slots_blank_and_offsets_exactly(self):
        layout = FORMATS_BY_CODE["L7160"]
        image = render_sheet([None, "Etikett 2\nÖl", "Etikett 3"], layout)
        scale = RENDER_DPI / 25.4
        x, y, w, h = layout.rect(0)
        blank = image.crop((round(x*scale), round(y*scale), round((x+w)*scale), round((y+h)*scale)))
        self.assertEqual(blank.getextrema(), ((255, 255),) * 3)
        self.assertIsNotNone(ImageChops.invert(image).getbbox())
        original = render_sheet(["Position"], layout)
        shifted = render_sheet(["Position"], layout, off_x=2, off_y=3)
        a, b = ImageChops.invert(original).getbbox(), ImageChops.invert(shifted).getbbox()
        self.assertLessEqual(abs(b[0] - a[0] - 2 * scale), 1)
        self.assertLessEqual(abs(b[1] - a[1] - 3 * scale), 1)


class CatalogTests(unittest.TestCase):
    def test_catalog_is_loaded_unique_and_curated_first(self):
        self.assertGreater(len(FORMATS), 500)
        self.assertEqual(FORMATS[:len(CURATED)], CURATED)
        self.assertEqual(len(FORMATS_BY_CODE), len(FORMATS))
        for code in ("3477", "3475", "3490", "4780", "3666", "3662", "L7651", "L7160"):
            self.assertIn(code, FORMATS_BY_CODE)
        self.assertTrue(FORMATS_BY_CODE["3662"].landscape)
        self.assertEqual(FORMATS_BY_CODE["4790"].shape, "round")
        f3477 = FORMATS_BY_CODE["3477"]
        self.assertEqual((f3477.width, f3477.height, f3477.per_sheet), (105, 41, 14))
        flexicom = search_formats("FX1574-W")
        self.assertEqual([f.code for f in flexicom], ["FX1574"])
        self.assertEqual(flexicom[0].per_sheet, 28)
        self.assertEqual(flexicom[0].rect(27), (105, 268.5, 90, 20))

    def test_search_matches_code_name_and_size(self):
        self.assertEqual(search_formats("3477")[0].code, "3477")
        self.assertTrue(all("ordner" in f.description.lower() for f in search_formats("Ordner")))
        self.assertIn("3475", [f.code for f in search_formats("70 × 36")])
        self.assertIn("L7160", [f.code for f in search_formats("63,5 38,1")])
        self.assertEqual(search_formats("gibtesnicht"), [])

    def test_landscape_round_and_rotated_rendering(self):
        scale = RENDER_DPI / 25.4
        spine = FORMATS_BY_CODE["3662"]
        self.assertGreater(render_sheet(["Ordner 2026"], spine).width,
                           render_sheet(["Ordner 2026"], spine).height)
        rotated = render_sheet(["Ordner 2026"], spine, rotate=True)
        left, top, width, height = spine.rect(0)
        box = rotated.crop((round(left * scale), round(top * scale),
                            round((left + width) * scale), round((top + height) * scale)))
        ink = ImageChops.invert(box).getbbox()
        self.assertGreater(ink[3] - ink[1], ink[2] - ink[0])  # text runs along the long side
        disc = FORMATS_BY_CODE["4790"]
        image = render_sheet(["Rund"], disc, border=True)
        x, y, _, _ = disc.rect(0)
        corner = image.getpixel((round((x + .3) * scale), round((y + .3) * scale)))
        self.assertEqual(corner, (255, 255, 255))  # nothing printed outside the circle


@unittest.skipUnless(os.name == "nt", "Windows GUI integration")
class AppTests(unittest.TestCase):
    def test_background_import_and_failure_clear_previous_data(self):
        import lager_label
        import time
        with tempfile.TemporaryDirectory() as folder, patch.object(lager_label, "SETTINGS_FILE", str(Path(folder) / "settings.json")), \
                patch.object(lager_label.App, "printer_list", return_value=[]):
            path = Path(folder) / "test.xlsm"
            book = Workbook()
            book.active["A1"] = "Imported label"
            book.save(path)
            book.close()
            app = lager_label.App()
            app.withdraw()
            try:
                tab = app.free_tab
                with patch("free_label_tab.filedialog.askopenfilename", return_value=str(path)), \
                        patch("free_label_tab.messagebox.showerror") as error:
                    for corrupt in (False, True):
                        if corrupt:
                            path.write_text("Broken Excel file")
                        tab.open_excel()
                        self.assertTrue(tab.loading)
                        self.assertEqual(str(tab.print_button.cget("state")), "disabled")
                        deadline = time.monotonic() + 10
                        while tab.loading and time.monotonic() < deadline:
                            app.update()
                            time.sleep(.01)
                        self.assertFalse(tab.loading)
                        if corrupt:
                            error.assert_called_once()
                            self.assertEqual(tab.items, [])
                            self.assertEqual(str(tab.print_button.cget("state")), "disabled")
                        else:
                            self.assertEqual(tab.items, ["Imported label"])
                            self.assertEqual(str(tab.print_button.cget("state")), "normal")
            finally:
                app.destroy()

    def test_import_options_print_routing_and_settings(self):
        import lager_label
        with tempfile.TemporaryDirectory() as folder, patch.object(lager_label, "SETTINGS_FILE", str(Path(folder) / "settings.json")), \
                patch.object(lager_label.App, "printer_list", return_value=[]):
            app = lager_label.App()
            app.withdraw()
            try:
                tab = app.free_tab
                self.assertEqual(len(app.notebook.tabs()), 3)
                tab.workbook = {"Test": [ExcelRow(1, "Header")] + [ExcelRow(i + 2, f"Label {i}") for i in range(25)]}
                tab.sheet_name.set("Test")
                tab.vars["skip_header"].set(True)
                tab.refresh()
                self.assertEqual(len(tab.sheets), 2)
                self.assertEqual(tab.sheets[0][0], "Label 0")
                self.assertEqual(tab.sheets[1][0], "Label 21")
                tab.turn_page(1)
                self.assertEqual(tab.page_index, 1)
                app.vars["printer"].set("Test printer")
                with patch("free_label_tab.print_excel_sheets") as printer, \
                        patch("free_label_tab.messagebox.showinfo"), patch.object(app, "devmode_for_print", return_value=None):
                    tab.do_print(test=True)
                    self.assertEqual(len(printer.call_args.args[1]), 1)
                    self.assertTrue(printer.call_args.args[3]["frames"])
                    self.assertEqual(printer.call_args.args[1][0][0], "Label 0")
                    tab.vars["start"].set("21")
                    tab.do_print()
                    self.assertEqual(len(printer.call_args.args[1]), 3)
                    self.assertEqual(printer.call_args.args[1][0][-1], "Label 0")
                    printer.reset_mock()
                    tab.vars["off_x"].set("NaN")
                    tab.do_print()
                    printer.assert_not_called()
                    self.assertEqual(str(tab.print_button.cget("state")), "disabled")
                tab.vars["off_x"].set("1,5")
                tab.refresh()
                stored = json.loads(Path(folder, "settings.json").read_text(encoding="utf-8"))
                self.assertEqual(stored["free_labels"]["off_x"], "1,5")
                self.assertTrue(stored["free_labels"]["skip_header"])
                self.assertNotIn("Label 0", json.dumps(stored))
                self.assertEqual(lager_label.build_sheets(["old"], 4), [[None, None, None, "old"]])
            finally:
                app.destroy()


@unittest.skipUnless(os.name == "nt", "Windows GDI integration")
class PrintTests(unittest.TestCase):
    def test_physical_page_mapping_and_cleanup(self):
        import win32con
        from unittest.mock import MagicMock
        from free_labels import print_excel_sheets
        caps = {win32con.LOGPIXELSX: 600, win32con.LOGPIXELSY: 600,
                win32con.PHYSICALOFFSETX: 42, win32con.PHYSICALOFFSETY: 60,
                win32con.PHYSICALWIDTH: round(210 / 25.4 * 600),
                win32con.PHYSICALHEIGHT: round(297 / 25.4 * 600)}
        for fails in (False, True):
            dc = MagicMock()
            dc.GetDeviceCaps.side_effect = caps.__getitem__
            dm = SimpleNamespace(Fields=0, Scale=50)
            with patch("win32print.OpenPrinter"), patch("win32print.ClosePrinter") as close, \
                    patch("win32print.GetPrinter", return_value={"pDevMode": dm}), \
                    patch("win32print.DocumentProperties"), patch("win32gui.CreateDC"), \
                    patch("win32ui.CreateDCFromHandle", return_value=dc), \
                    patch("PIL.ImageWin.Dib") as dib:
                if fails:
                    dib.return_value.draw.side_effect = RuntimeError("driver failed")
                    with self.assertRaisesRegex(RuntimeError, "driver failed"):
                        print_excel_sheets("Printer", [["A"]], FORMATS[0], {})
                    dc.AbortDoc.assert_called_once()
                else:
                    print_excel_sheets("Printer", [["A"], ["B"]], FORMATS[0], {})
                    self.assertEqual(dc.StartPage.call_count, 2)
                    self.assertEqual(dc.EndPage.call_count, 2)
                    dc.EndDoc.assert_called_once()
                    dc.AbortDoc.assert_not_called()
                    self.assertEqual(dib.return_value.draw.call_args.args[1],
                        (-42, -60, caps[win32con.PHYSICALWIDTH] - 42, caps[win32con.PHYSICALHEIGHT] - 60))
                self.assertEqual(dm.PaperSize, win32con.DMPAPER_A4)
                self.assertEqual(dm.Orientation, win32con.DMORIENT_PORTRAIT)
                self.assertEqual(dm.Scale, 100)
                close.assert_called_once()
                dc.DeleteDC.assert_called_once()


if __name__ == "__main__":
    unittest.main()
